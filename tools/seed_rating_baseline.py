#!/usr/bin/env python3
"""Write the recorded rating snapshots into the workbook's Rating_Tracker tab.

data/ratings.csv is long (one row per entity x platform x week); the sheet is
wide (one row per entity per week). This pivots between them so the sheet ships
carrying the baseline instead of asking someone to retype seven numbers and
risk a transcription error against the series everything is measured from.
"""
import collections
import datetime as dt
import os
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "scripts"))
import hrintel as H

import openpyxl
from openpyxl.styles import Font

BOOK = sys.argv[1] if len(sys.argv) > 1 else "templates/HR_Intelligence_Master_Tracker.xlsx"

rows = H.read_csv(H.RATINGS_CSV)
if not rows:
    print("No rating snapshots recorded; nothing to seed.")
    raise SystemExit(0)

names = H.entity_names()
wide = collections.OrderedDict()
for r in rows:
    key = (r["week_of"], r["entity"])
    entry = wide.setdefault(key, {"notes": [], "checked": r.get("captured_at", ""),
                                  "captured_by": r.get("captured_by", "")})
    entry[r["platform"]] = r
    if r.get("notes"):
        entry["notes"].append(f"{r['platform']}: {r['notes']}")

wb = openpyxl.load_workbook(BOOK)
ws = wb["Rating_Tracker"]
font = Font(name="Arial")

# Write by header name, not by column position. The positional version broke
# silently the first time a column was inserted: the numbers still landed, one
# column to the left of where they belonged.
cfg = H.load_yaml("column_map")
spec = cfg["tabs"]["ratings"]


def norm(name) -> str:
    return "".join(ch for ch in str(name or "").lower() if ch.isalnum())


header = {norm(c.value): i for i, c in enumerate(ws[1], start=1) if c.value}


def column_for(spellings) -> int | None:
    for spelling in spellings:
        index = header.get(norm(spelling))
        if index:
            return index
    return None


COUNTS = {"review_count", "recommend_pct", "ceo_approval_pct"}
DATES = {"week_of", "captured_at"}
TEXT = {"entity", "captured_by", "url", "notes"}


def cast(field: str, raw):
    text = str(raw or "").strip()
    if not text:
        return None
    if field in DATES:
        parsed = H.parse_date(text)
        return dt.date.fromisoformat(parsed.isoformat()) if parsed else text
    if field in COUNTS:
        # Not 'or None': Westbury's Glassdoor recommend rate is 0%, which is
        # the most pointed number on the tab and must not render as blank.
        return H.to_int(text, None)
    if field in TEXT:
        return text
    return H.to_float(text)


missing: set[str] = set()
for i, ((week, entity), entry) in enumerate(sorted(wide.items()), start=2):
    cells: dict[int, object] = {}

    flat = {"week_of": week, "entity": names.get(entity, entity),
            "captured_at": entry.get("checked", ""),
            "captured_by": entry.get("captured_by", ""),
            "notes": ""}
    for field, spellings in spec["columns"].items():
        index = column_for(spellings)
        if index is None:
            missing.add(field)
            continue
        value = flat.get(field, "")
        cells[index] = (names.get(entity, entity) if field == "entity"
                        else cast(field, value))

    for platform, block in (spec.get("per_platform") or {}).items():
        source = entry.get(platform, {})
        for field, spellings in block.items():
            index = column_for(spellings)
            if index is None:
                missing.add(f"{platform}.{field}")
                continue
            cells[index] = cast(field, source.get(field))

    for index, value in cells.items():
        cell = ws.cell(row=i, column=index, value=value)
        cell.font = font

wb.save(BOOK)
print(f"Seeded {len(wide)} Rating_Tracker row(s) into {BOOK}")
for (week, entity) in sorted(wide):
    print(f"  {week}  {names.get(entity, entity)}")
if missing:
    print("\nNo column in the sheet for: " + ", ".join(sorted(missing)))
