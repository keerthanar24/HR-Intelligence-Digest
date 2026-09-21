"""Fill Keyword_Matrix from the alias register, and add the Guide tab."""
import subprocess, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

REPO = "/home/user/HR-Intelligence-Digest"
DST = sys.argv[1]
FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="D9E2EC")
HDR_FONT = Font(name=FONT, bold=True, color="1F2933")
BODY = Font(name=FONT)
H2 = Font(name=FONT, bold=True, size=12, color="1F2933")
NOTE = Font(name=FONT, italic=True, color="52606D")

wb = openpyxl.load_workbook(DST)

# ---- Keyword_Matrix, generated from config/entities.yaml ----
out = subprocess.run([sys.executable, "scripts/alert_queries.py", "--format", "matrix"],
                     cwd=REPO, capture_output=True, text=True, check=True).stdout
rows = [line.split("\t") for line in out.splitlines() if line.count("\t") == 3]

ws = wb["Keyword_Matrix"]
for row in ws.iter_rows(min_row=1, max_row=max(ws.max_row, 1), max_col=max(ws.max_column, 1)):
    for cell in row:
        cell.value = None

widths = [22, 62, 70, 70]
for i, (head, width) in enumerate(zip(rows[0], widths), start=1):
    c = ws.cell(row=1, column=i, value=head)
    c.font, c.fill = HDR_FONT, HDR_FILL
    c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(i)].width = width
ws.row_dimensions[1].height = 30
ws.freeze_panes = "A2"

for r, row in enumerate(rows[1:], start=2):
    for i, value in enumerate(row, start=1):
        c = ws.cell(row=r, column=i, value=value)
        c.font = BODY
        c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[r].height = 58

note_row = len(rows) + 2
ws.cell(row=note_row, column=1,
        value="Generated from config/entities.yaml by: python3 scripts/alert_queries.py "
              "--format matrix").font = NOTE
ws.cell(row=note_row + 1, column=1,
        value="X-ray strings target review and discussion sites only. Profile X-ray "
              "(site:linkedin.com/in/) is excluded on purpose — enumerating individuals' "
              "profiles is outside the brief.").font = NOTE

# ---- Guide ----
if "Guide" in wb.sheetnames:
    del wb["Guide"]
g = wb.create_sheet("Guide", 0)
g.column_dimensions["A"].width = 26
g.column_dimensions["B"].width = 104

LINES = [
    ("h", "HR Intelligence Digest — Tracker Guide"),
    ("n", "60-day awareness trial. First-level awareness only: no action items or HR process "
          "changes follow from a digest. Red Flags are the one exception and escalate same-day."),
    ("b", ""),
    ("h", "Scope — printed at the foot of every digest"),
    ("kv", ("In scope", "Public, employment-related commentary: culture, management, pay, "
                        "appraisals, exits, layoffs, work hours, interview experiences, "
                        "onboarding, harassment or safety allegations.")),
    ("kv", ("Out of scope", "Customer and seller complaints, product reviews, and any "
                            "surveillance of an individual's personal social media. We do not "
                            "open private or connection-gated content, and we do not try to "
                            "identify anonymous reviewers.")),
    ("b", ""),
    ("h", "Tabs"),
    ("kv", ("Raw_Data_Log", "One row per public item found. Week Of fills itself from Date Posted.")),
    ("kv", ("Rating_Tracker", "One row per entity per week. APPEND a new row each week — "
                              "overwriting destroys the history the digest compares against.")),
    ("kv", ("Escalations", "One row per red flag. The only tab that may carry an accused "
                           "individual's name. Move it to a separate file if this sheet is "
                           "ever shared beyond the four recipients.")),
    ("kv", ("Keyword_Matrix", "Search strings, generated from the alias register.")),
    ("b", ""),
    ("h", "Sentiment"),
    ("kv", ("Positive / Negative", "A clear view on something specific.")),
    ("kv", ("Neutral", "Nothing evaluative said. Most news items.")),
    ("kv", ("Mixed", "Real praise AND a real complaint in one post — 'great team, terrible pay'.")),
    ("n", "Tag the post, not your view of the company. Star ratings do not decide the tag. "
          "If you cannot choose in ten seconds, pick Mixed and move on. Tag a whole week in "
          "one sitting — drift within a week is what makes week-on-week comparison "
          "meaningless."),
    ("b", ""),
    ("h", "Topic Category"),
    ("kv", ("Compensation", "Paid too little — pay level, benefits, market comparison.")),
    ("kv", ("Payroll Delay", "NOT paid — late or withheld salary, full-and-final, PF/ESIC. "
                             "This one is a red-flag trigger; Compensation is not.")),
    ("kv", ("Appraisals", "Appraisal cycle, increments, ratings, promotions.")),
    ("kv", ("Exits / Layoffs", "Resignations and notice period vs. terminations and restructuring.")),
    ("b", ""),
    ("h", "Red Flags — same day, do not wait for Monday"),
    ("kv", ("Names Individual", "A named person is accused of specific conduct.")),
    ("kv", ("Harassment / Safety", "Harassment, discrimination, retaliation, unsafe conditions.")),
    ("kv", ("Non-payment", "Unpaid or withheld salary, FnF, PF/ESIC or statutory dues.")),
    ("kv", ("Legal / Regulatory", "Labour commissioner, tribunal, police complaint, legal notice.")),
    ("kv", ("Public Escalation Risk", "Traction, media pickup, a thread gaining momentum.")),
    ("n", "A bad review is not a red flag — bad reviews are the normal content of this "
          "digest. Set Red Flag Escalated = Yes, add the Reason, log a row on Escalations, "
          "and send the alert the same day. The alert never carries the accused person's "
          "name; it links to the source instead."),
    ("b", ""),
    ("h", "Weekly routine"),
    ("kv", ("1. Sweep", "Open the 12 company pages. Log every new item. Record the rating and "
                        "review count even when nothing changed.")),
    ("kv", ("2. Tag", "Fill Summary Overview, Sentiment, Topic, Status for every Needs Review row.")),
    ("kv", ("3. Red flags", "Confirm, log, send — same day.")),
    ("kv", ("4. Build", "Download as .xlsx, then: python3 scripts/import_sheet.py <file> "
                        "and python3 scripts/build_digest.py")),
    ("b", ""),
    ("h", "Colour"),
    ("kv", ("Red row", "Red Flag Escalated = Yes")),
    ("kv", ("Amber row", "Status = Needs Review. A row still amber on send day has not been tagged.")),
]

r = 1
for kind, payload in LINES:
    if kind == "h":
        c = g.cell(row=r, column=1, value=payload); c.font = H2
        g.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    elif kind == "n":
        c = g.cell(row=r, column=1, value=payload); c.font = NOTE
        c.alignment = Alignment(wrap_text=True, vertical="top")
        g.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        g.row_dimensions[r].height = 44
    elif kind == "kv":
        k, v = payload
        a = g.cell(row=r, column=1, value=k); a.font = Font(name=FONT, bold=True)
        a.alignment = Alignment(vertical="top")
        b = g.cell(row=r, column=2, value=v); b.font = BODY
        b.alignment = Alignment(wrap_text=True, vertical="top")
        g.row_dimensions[r].height = 30 if len(v) > 90 else 15
    r += 1

g.sheet_view.showGridLines = False
wb.save(DST)
print("populated:", DST)
