#!/usr/bin/env python3
"""Import the HR Intelligence Master Tracker workbook into the canonical CSVs.

The sheet keeps its own friendly column names and dropdowns; this translates
them using config/column_map.yaml so the digest scripts can read it. Nothing is
silently dropped — every unmapped column and unrecognised value is reported.

    python3 scripts/import_sheet.py tracker.xlsx            # whole workbook
    python3 scripts/import_sheet.py --dry-run tracker.xlsx  # report only
    python3 scripts/import_sheet.py --csv Raw_Data_Log.csv --tab mentions

Download the workbook as .xlsx (File > Download > Microsoft Excel) and point
this at it — that is one step instead of exporting three tabs separately.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import hrintel as H  # noqa: E402

RED_FLAG_SENTINEL = "__red_flag__"

TARGETS = {
    "mentions": (H.MENTIONS_CSV, H.MENTION_FIELDS),
    "ratings": (H.RATINGS_CSV, H.RATING_FIELDS),
    "escalations": (H.ESCALATIONS_CSV, H.ESCALATION_FIELDS),
}


def key(name: str) -> str:
    """Header/value comparison key: case, spaces and punctuation insensitive."""
    return "".join(ch for ch in str(name or "").lower() if ch.isalnum())


class Notes:
    """Collects what the import could not do, so it is reported rather than lost."""

    def __init__(self) -> None:
        self.warnings: list[str] = []
        self.seen: set[str] = set()

    def warn(self, message: str, once: bool = False) -> None:
        if once:
            if message in self.seen:
                return
            self.seen.add(message)
        self.warnings.append(message)


def build_lookup(columns: dict) -> dict[str, str]:
    """{normalised sheet header: canonical field}"""
    lookup = {}
    for field, spellings in columns.items():
        for spelling in spellings:
            lookup[key(spelling)] = field
    return lookup


def map_value(raw, kind: str, values: dict, notes: Notes, where: str):
    """Translate a sheet value to the canonical vocabulary."""
    raw = ("" if raw is None else str(raw)).strip()
    if not raw:
        return ""
    table = {key(k): v for k, v in (values.get(kind) or {}).items()}
    mapped = table.get(key(raw))
    if mapped is None:
        notes.warn(f"{where}: {kind} value {raw!r} is not in config/column_map.yaml "
                   f"— add it there or fix the dropdown", once=True)
        return raw
    return mapped


def as_date(raw) -> str:
    if raw in (None, ""):
        return ""
    if isinstance(raw, dt.datetime):
        return raw.date().isoformat()
    if isinstance(raw, dt.date):
        return raw.isoformat()
    parsed = H.parse_date(str(raw))
    return parsed.isoformat() if parsed else ""


def read_xlsx(path: str) -> dict[str, list[list]]:
    try:
        import openpyxl
    except ImportError:
        raise SystemExit(
            "Reading .xlsx needs openpyxl:  python3 -m pip install openpyxl\n"
            "Or export each tab as CSV and use --csv."
        )
    workbook = openpyxl.load_workbook(path, data_only=True)
    return {
        name: [list(row) for row in workbook[name].iter_rows(values_only=True)]
        for name in workbook.sheetnames
    }


def read_csv_file(path: str) -> list[list]:
    with open(path, newline="", encoding="utf-8-sig") as fh:
        return [row for row in csv.reader(fh)]


def to_records(rows: list[list], lookup: dict[str, str]) -> tuple[list[dict], list[str]]:
    """Header row + data rows -> dicts keyed by canonical field name."""
    if not rows:
        return [], []
    header = rows[0]
    mapping, unmapped = {}, []
    for index, name in enumerate(header):
        if name is None or not str(name).strip():
            continue
        field = lookup.get(key(name))
        if field:
            mapping[index] = field
        else:
            unmapped.append(str(name).strip())

    records = []
    for row in rows[1:]:
        record = {}
        for index, field in mapping.items():
            if index < len(row):
                record[field] = row[index]
        # Skip rows that are entirely blank, which spreadsheets have in bulk.
        if any(str(v).strip() for v in record.values() if v is not None):
            records.append(record)
    return records, unmapped


def import_mentions(rows, cfg, notes: Notes) -> list[dict]:
    lookup = build_lookup(cfg["tabs"]["mentions"]["columns"])
    records, unmapped = to_records(rows, lookup)
    for name in unmapped:
        notes.warn(f"Raw_Data_Log: column {name!r} is not mapped and was skipped")

    out: list[dict] = []
    for number, record in enumerate(records, start=2):
        where = f"Raw_Data_Log row {number}"
        row = {field: "" for field in H.MENTION_FIELDS}
        row.update({k: ("" if v is None else str(v).strip()) for k, v in record.items()})

        row["captured_at"] = as_date(record.get("captured_at")) or dt.date.today().isoformat()
        row["post_date"] = as_date(record.get("post_date"))

        posted = H.parse_date(row["post_date"])
        recorded = H.parse_date(as_date(record.get("week_of")))
        if recorded:
            row["week_of"] = H.monday_of(recorded).isoformat()
        elif posted:
            row["week_of"] = H.monday_of(posted).isoformat()
        else:
            captured = H.parse_date(row["captured_at"])
            row["week_of"] = H.monday_of(captured or dt.date.today()).isoformat()
            notes.warn(f"{where}: no Date Posted — filed by Date Logged instead, "
                       "which can put it in the wrong week")

        row["entity"] = map_value(record.get("entity"), "entity", cfg["values"], notes, where)
        row["platform"] = map_value(record.get("platform"), "platform", cfg["values"], notes, where)
        row["red_flag"] = map_value(record.get("red_flag"), "red_flag", cfg["values"], notes, where)
        for field in ("author_type", "red_flag_reason", "status"):
            if record.get(field):
                row[field] = map_value(record.get(field), field, cfg["values"], notes, where)

        sentiment = map_value(record.get("sentiment"), "sentiment", cfg["values"], notes, where)
        if sentiment == RED_FLAG_SENTINEL:
            # 'Red Flag' is not a point on a sentiment scale. Set the flag and
            # leave sentiment blank rather than invent a score; the digest then
            # reports the row as untagged instead of quietly skewing the mean.
            row["red_flag"] = "yes"
            row["sentiment"] = ""
            notes.warn(f"{where}: Sentiment was 'Red Flag'. Imported as red_flag=yes with no "
                       "sentiment — re-tag it Positive/Neutral/Negative so it counts in "
                       "section 1", once=False)
        else:
            row["sentiment"] = sentiment

        themes = record.get("themes")
        if themes:
            mapped = [
                map_value(part, "themes", cfg["values"], notes, where)
                for part in str(themes).replace(";", "|").replace(",", "|").split("|")
                if str(part).strip()
            ]
            row["themes"] = "|".join(t for t in mapped if t)

        if H.is_yes(row["red_flag"]) and not row.get("red_flag_reason"):
            notes.warn(f"{where}: flagged red but the sheet has no Red Flag Reason column — "
                       "section 5 cannot say which of the five triggers fired", once=True)

        if not row.get("status"):
            row["status"] = "reviewed" if row["sentiment"] else "needs_review"
        if not row.get("names_individual"):
            row["names_individual"] = ""
        if not row.get("captured_by"):
            row["captured_by"] = "sheet"
        out.append(row)

    # Assign ids only where the sheet has no Mention ID column of its own.
    used: dict[str, int] = {}
    for row in out:
        if row.get("mention_id"):
            continue
        week = H.parse_date(row["week_of"]) or dt.date.today()
        prefix = f"M-{week.strftime('%Y%m%d')}-"
        used[prefix] = used.get(prefix, 0) + 1
        row["mention_id"] = f"{prefix}{used[prefix]:03d}"
    return out


def import_ratings(rows, cfg, notes: Notes) -> list[dict]:
    spec = cfg["tabs"]["ratings"]
    lookup = build_lookup(spec["columns"])
    per_platform = spec.get("per_platform", {})
    for platform, fields in per_platform.items():
        for field, spellings in fields.items():
            for spelling in spellings:
                lookup[key(spelling)] = f"{platform}::{field}"

    ignored = {key(name) for name in spec.get("ignore", [])}
    records, unmapped = to_records(rows, lookup)
    for name in unmapped:
        if key(name) in ignored:
            continue
        notes.warn(f"Rating_Tracker: column {name!r} is not mapped and was skipped")

    if any(key(n) == key("Total Reviews Count") for n in unmapped):
        notes.warn(
            "Rating_Tracker: 'Total Reviews Count' is one column serving two platforms that "
            "have separate ratings, so it cannot be attributed and was not imported. Split it "
            "into 'AmbitionBox Reviews Count' and 'Glassdoor Reviews Count' and it imports "
            "automatically."
        )

    out = []
    for number, record in enumerate(records, start=2):
        where = f"Rating_Tracker row {number}"
        entity = map_value(record.get("entity"), "entity", cfg["values"], notes, where)
        captured = as_date(record.get("captured_at")) or dt.date.today().isoformat()
        recorded = H.parse_date(as_date(record.get("week_of")))
        week = H.monday_of(recorded or H.parse_date(captured) or dt.date.today())

        for platform in per_platform:
            values = {
                field.split("::", 1)[1]: value
                for field, value in record.items()
                if field.startswith(f"{platform}::")
            }
            if not any(str(v).strip() for v in values.values() if v is not None):
                continue
            row = {field: "" for field in H.RATING_FIELDS}
            row.update(
                {
                    "week_of": week.isoformat(),
                    "captured_at": captured,
                    "captured_by": str(record.get("captured_by") or "sheet"),
                    "entity": entity,
                    "platform": platform,
                    "overall_rating": str(values.get("overall_rating") or "").strip(),
                    "review_count": str(values.get("review_count") or "").strip(),
                    "url": str(values.get("url") or "").strip(),
                    "notes": str(record.get("notes") or "").strip(),
                }
            )
            out.append(row)

    if out:
        weeks = {row["week_of"] for row in out}
        if len(weeks) == 1:
            notes.warn(
                f"Rating_Tracker holds one snapshot ({sorted(weeks)[0]}). Section 2 reports "
                "movement between consecutive weekly snapshots, so APPEND a new row each week "
                "rather than overwriting — an overwritten row destroys the series permanently."
            )
    return out


def import_escalations(rows, cfg, notes: Notes) -> list[dict]:
    lookup = build_lookup(cfg["tabs"]["escalations"]["columns"])
    records, unmapped = to_records(rows, lookup)
    for name in unmapped:
        notes.warn(f"Escalations: column {name!r} is not mapped and was skipped")
    out = []
    for record in records:
        row = {field: "" for field in H.ESCALATION_FIELDS}
        row.update({k: ("" if v is None else str(v).strip()) for k, v in record.items()})
        where = f"Escalations row {record.get('escalation_id') or '?'}"
        if record.get("severity"):
            row["severity"] = map_value(record["severity"], "severity", cfg["values"], notes, where)
        if record.get("reason"):
            row["reason"] = map_value(record["reason"], "red_flag_reason", cfg["values"], notes, where)
        if record.get("status"):
            row["status"] = map_value(record["status"], "escalation_status", cfg["values"], notes, where)
        row["raised_at"] = as_date(record.get("raised_at"))
        row["notified_at"] = as_date(record.get("notified_at"))
        row["closed_at"] = as_date(record.get("closed_at"))
        out.append(row)
    return out


def find_tab(sheets: dict, names: list[str]) -> str | None:
    wanted = {key(n) for n in names}
    for name in sheets:
        if key(name) in wanted:
            return name
    return None


def write(path: str, fields: list[str], rows: list[dict]) -> None:
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({f: row.get(f, "") for f in fields})


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", nargs="?", help="path to the .xlsx workbook")
    parser.add_argument("--csv", help="import a single exported CSV instead")
    parser.add_argument("--tab", choices=sorted(TARGETS),
                        help="which schema the --csv file follows")
    parser.add_argument("--dry-run", action="store_true", help="report without writing")
    args = parser.parse_args()

    if not args.workbook and not args.csv:
        parser.error("give a workbook path, or --csv with --tab")
    if args.csv and not args.tab:
        parser.error("--csv needs --tab to say which schema it follows")

    cfg = H.load_yaml("column_map")
    notes = Notes()
    importers = {
        "mentions": import_mentions,
        "ratings": import_ratings,
        "escalations": import_escalations,
    }

    if args.csv:
        sheets = {args.tab: read_csv_file(args.csv)}
        wanted = [args.tab]
    else:
        sheets = read_xlsx(args.workbook)
        print(f"Workbook tabs: {', '.join(sheets)}\n")
        wanted = list(TARGETS)

    results: dict[str, list[dict]] = {}
    for tab in wanted:
        if args.csv:
            rows = sheets[tab]
        else:
            found = find_tab(sheets, cfg["tabs"][tab]["sheet_names"])
            if not found:
                expected = cfg["tabs"][tab]["sheet_names"][0]
                notes.warn(f"No '{expected}' tab in the workbook — "
                           f"{tab}.csv left unchanged")
                continue
            rows = sheets[found]
        results[tab] = importers[tab](rows, cfg, notes)

    for tab, rows in results.items():
        path, fields = TARGETS[tab]
        print(f"  {tab}: {len(rows)} row(s)"
              f"{' (not written — dry run)' if args.dry_run else f' -> {os.path.relpath(path, H.ROOT)}'}")
        if not args.dry_run:
            write(path, fields, rows)

    if notes.warnings:
        print(f"\n{len(notes.warnings)} note(s):\n")
        for message in notes.warnings:
            print(f"  - {message}")

    if not args.dry_run and results:
        print("\nNext: python3 scripts/validate_data.py --week <monday>")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
